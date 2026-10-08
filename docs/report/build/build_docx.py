"""Builds the report as DOCX (python-docx). One source, two flavours:
  - "pdf":  Vazirmatn font and a static table of contents (rendered to the 54-page PDF by LibreOffice)
  - "word": B Nazanin font and a Word table-of-contents field that Word refreshes on opening
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).parent))
from content_front import *  # noqa
from content_ch12 import CH1, CH2  # noqa
from content_ch3 import CH3  # noqa
from content_ch4 import CH4  # noqa
from content_ch567 import CH5, CH6, CH7, REFS, APPENDIX  # noqa

FIG_DIR = Path(__file__).resolve().parent.parent / "figures"
NAVY, NAVY_LIGHT, NAVY_PALE, INK, MUTED = "1F3A6E", "E3EAF6", "F3F6FB", RGBColor(0x14, 0x21, 0x3D), RGBColor(0x56, 0x66, 0x86)
DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
fa = lambda n: str(n).translate(DIGITS)

CONFIG = {
    "pdf": dict(font="Vazirmatn UI FD", latin="Vazirmatn UI FD", body=10.5, table=9, small=8.5, line=1.4, h1=19, h2=13.5, h3=11.5, code_font="DejaVu Sans Mono"),
    "word": dict(font="B Nazanin", latin="Times New Roman", body=13, table=11, small=10, line=1.35, h1=20, h2=16, h3=14, code_font="Courier New"),
}


class Report:
    def __init__(self, flavour, toc_pages=None, list_pages=None, overrides=None):
        self.cfg, self.flavour = {**CONFIG[flavour], **(overrides or {})}, flavour
        self.toc_pages, self.list_pages = toc_pages or {}, list_pages or {}
        self.doc = Document()
        self.chapter, self.fig_no, self.tbl_no = 0, 0, 0
        self.headings, self.figures, self.tables = [], [], []
        self.setup_styles()

    # ---------------------------------------------------------------- low level helpers
    def font(self, run, size=None, bold=False, color=None, mono=False, italic=False):
        name = self.cfg["code_font"] if mono else self.cfg["font"]
        latin = self.cfg["code_font"] if mono else self.cfg["latin"]
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
        for attr, value in (("w:ascii", latin), ("w:hAnsi", latin), ("w:cs", name), ("w:eastAsia", name)):
            fonts.set(qn(attr), value)
        if size:
            run.font.size = Pt(size)
            szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rpr.append(szcs)
        if bold:
            run.font.bold = True
            bcs = OxmlElement("w:bCs"); rpr.append(bcs)
        if italic:
            run.font.italic = True
        if color is not None:
            run.font.color.rgb = color if isinstance(color, RGBColor) else RGBColor.from_string(color)
        if not mono:
            rpr.append(OxmlElement("w:rtl"))

    def setup_styles(self):
        normal = self.doc.styles["Normal"]
        normal.font.name = self.cfg["latin"]
        normal.font.size = Pt(self.cfg["body"])
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
        for attr, value in (("w:ascii", self.cfg["latin"]), ("w:hAnsi", self.cfg["latin"]), ("w:cs", self.cfg["font"]), ("w:eastAsia", self.cfg["font"])):
            fonts.set(qn(attr), value)
        szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(self.cfg["body"] * 2))); rpr.append(szcs)
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.3)
        # Heading styles: real heading levels so Word can build its own contents list.
        for level, size in ((1, self.cfg["h1"]), (2, self.cfg["h2"]), (3, self.cfg["h3"])):
            style = self.doc.styles[f"Heading {level}"]
            style.font.name = self.cfg["latin"]; style.font.size = Pt(size); style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(NAVY if level < 3 else "172D5C")
            srpr = style.element.get_or_add_rPr()
            sf = srpr.find(qn("w:rFonts"))
            if sf is None:
                sf = OxmlElement("w:rFonts"); srpr.insert(0, sf)
            for attr, value in (("w:ascii", self.cfg["latin"]), ("w:hAnsi", self.cfg["latin"]), ("w:cs", self.cfg["font"]), ("w:eastAsia", self.cfg["font"])):
                sf.set(qn(attr), value)
            sz = OxmlElement("w:szCs"); sz.set(qn("w:val"), str(int(size * 2))); srpr.append(sz)
            style.paragraph_format.keep_with_next = True

    def rtl(self, paragraph, align="both"):
        ppr = paragraph._p.get_or_add_pPr()
        ppr.append(OxmlElement("w:bidi"))
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), align); ppr.append(jc)

    def spacing(self, paragraph, before=0, after=6, line=None):
        pf = paragraph.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = line or self.cfg["line"]

    def shade(self, cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); tcpr.append(shd)

    def borders(self, table, color="C9D4EA", size=4):
        tblpr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color); borders.append(e)
        tblpr.append(borders)

    def field(self, paragraph, instruction, placeholder="۱", size=None, color=None):
        def run_with(child):
            r = paragraph.add_run(); r._r.append(child); return r
        begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" {instruction} "
        sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
        end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
        run_with(begin); run_with(instr); run_with(sep)
        shown = paragraph.add_run(placeholder); self.font(shown, size, color=color)
        run_with(end)

    # ---------------------------------------------------------------- block renderers
    def para(self, text, size=None, bold=False, color=None, align="both", before=0, after=6, line=None, keep=False, style=None):
        paragraph = self.doc.add_paragraph(style=style)
        self.rtl(paragraph, align); self.spacing(paragraph, before, after, line)
        run = paragraph.add_run(text); self.font(run, size or self.cfg["body"], bold, color)
        if keep:
            paragraph.paragraph_format.keep_with_next = True
        return paragraph

    def heading(self, level, text):
        if level == 1:
            if text.startswith("فصل"):
                self.chapter += 1
            self.page_break_before = True
        paragraph = self.doc.add_paragraph(style=f"Heading {level}")
        self.rtl(paragraph, "right" if False else "both")
        if level == 1:
            paragraph.paragraph_format.page_break_before = True
        size = {1: self.cfg["h1"], 2: self.cfg["h2"], 3: self.cfg["h3"]}[level]
        self.spacing(paragraph, {1: 0, 2: 14, 3: 10}[level], {1: 16, 2: 6, 3: 4}[level], 1.2)
        run = paragraph.add_run(text); self.font(run, size, True, NAVY if level < 3 else "172D5C")
        if level == 1:  # accent line below chapter titles
            ppr = paragraph._p.get_or_add_pPr()
            border = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            for k, v in (("val", "single"), ("sz", "12"), ("space", "6"), ("color", "FF6B4A")):
                bottom.set(qn(f"w:{k}"), v)
            border.append(bottom); ppr.append(border)
        if level <= 2:
            self.headings.append((level, text))

    def bullets(self, items, numbered=False):
        for index, item in enumerate(items, start=1):
            paragraph = self.doc.add_paragraph()
            self.rtl(paragraph); self.spacing(paragraph, 0, 3)
            paragraph.paragraph_format.right_indent = Cm(0.9)
            paragraph.paragraph_format.first_line_indent = Cm(-0.6)
            mark = f"{fa(index)}. " if numbered else "• "
            self.font(paragraph.add_run(mark), self.cfg["body"], True, "FF6B4A" if not numbered else NAVY)
            self.font(paragraph.add_run(item), self.cfg["body"])

    def caption(self, kind, text, above):
        if kind == "شکل":
            self.fig_no += 1; number = f"{fa(self.chapter)}-{fa(self.fig_no)}"; self.figures.append((number, text))
        else:
            self.tbl_no += 1; number = f"{fa(self.chapter)}-{fa(self.tbl_no)}"; self.tables.append((number, text))
        paragraph = self.doc.add_paragraph()
        self.rtl(paragraph, "center"); self.spacing(paragraph, 4 if above else 2, 4 if above else 10, 1.2)
        paragraph.paragraph_format.keep_with_next = above
        outline = OxmlElement("w:outlineLvl"); outline.set(qn("w:val"), "3"); paragraph._p.get_or_add_pPr().append(outline)  # bookmark for page lookup
        self.font(paragraph.add_run(f"{kind} {number}: "), self.cfg["small"] + 1, True, NAVY)
        self.font(paragraph.add_run(text), self.cfg["small"] + 1, False, MUTED)

    def figure(self, name, caption, width_cm, max_h=11.5):
        path = FIG_DIR / name
        with Image.open(path) as image:
            ratio = image.height / image.width
        width = min(width_cm, max_h / ratio)
        paragraph = self.doc.add_paragraph()
        self.rtl(paragraph, "center"); self.spacing(paragraph, 6, 2, 1.0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Cm(width))
        self.caption("شکل", caption, above=False)

    def table(self, caption, headers, rows, widths, small):
        self.caption("جدول", caption, above=True)
        size = self.cfg["small"] if small else self.cfg["table"]
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = table._tbl.tblPr
        tblpr.append(OxmlElement("w:bidiVisual"))
        layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblpr.append(layout)
        self.borders(table)
        total = sum(widths) if widths else 0
        widths = widths or [16.0 / len(headers)] * len(headers)
        scale = 16.0 / sum(widths)
        for r_index, row in enumerate([headers] + rows):
            for c_index, value in enumerate(row):
                cell = table.cell(r_index, c_index)
                cell.width = Cm(widths[c_index] * scale)
                cell.text = ""
                paragraph = cell.paragraphs[0]
                self.rtl(paragraph, "right" if False else "both"); self.spacing(paragraph, 1, 1, 1.15)
                latin_only = bool(re.fullmatch(r"[\x00-\x7F]+", str(value))) and r_index > 0
                if latin_only:
                    paragraph.alignment = None
                run = paragraph.add_run(str(value))
                if r_index == 0:
                    self.font(run, size, True, "FFFFFF"); self.shade(cell, NAVY)
                else:
                    self.font(run, size, False, None)
                    if r_index % 2 == 0:
                        self.shade(cell, NAVY_PALE)
        header_row = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true"); header_row.append(header)
        for row in table.rows:  # never split a row across pages
            trpr = row._tr.get_or_add_trPr(); cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true"); trpr.append(cant)
        spacer = self.doc.add_paragraph(); self.spacing(spacer, 0, 6, 1.0)

    def code(self, text, caption):
        if caption:
            paragraph = self.para(caption, self.cfg["small"], True, NAVY, "center", 4, 2, 1.1, keep=True)
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self.borders(table, "C9D4EA", 4)
        cell = table.cell(0, 0); cell.width = Cm(16); self.shade(cell, "F3F6FB")
        cell.text = ""
        lines = text.split("\n")
        for index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            self.spacing(paragraph, 0, 0, 1.15)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(line if line else " "); self.font(run, 8.5 if self.flavour == "pdf" else 9, False, INK, mono=True)
        spacer = self.doc.add_paragraph(); self.spacing(spacer, 0, 6, 1.0)

    def callout(self, text):
        table = self.doc.add_table(rows=1, cols=1); self.borders(table, "FF6B4A", 6)
        cell = table.cell(0, 0); self.shade(cell, "FFF1ED"); cell.text = ""
        paragraph = cell.paragraphs[0]; self.rtl(paragraph); self.spacing(paragraph, 2, 2, 1.25)
        self.font(paragraph.add_run(text), self.cfg["body"] - 0.5)
        spacer = self.doc.add_paragraph(); self.spacing(spacer, 0, 6, 1.0)

    def blocks(self, blocks):
        for block in blocks:
            kind = block[0]
            if kind in ("h1", "h2", "h3"):
                self.heading(int(kind[1]), block[1])
            elif kind == "p":
                self.para(block[1])
            elif kind == "ul":
                self.bullets(block[1])
            elif kind == "ol":
                self.bullets(block[1], True)
            elif kind == "fig":
                self.figure(block[1], block[2], block[3])
            elif kind == "table":
                self.table(block[1], block[2], block[3], block[4], block[5])
            elif kind == "code":
                self.code(block[1], block[2])
            elif kind == "note":
                self.callout(block[1])

    # ---------------------------------------------------------------- front matter and sections
    def new_section(self, restart=None, header_text=None, footer_numbers=False):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False; section.footer.is_linked_to_previous = False
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                paragraph.text = ""
        if restart:
            pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), str(restart)); section._sectPr.append(pg)
        if header_text:
            paragraph = section.header.paragraphs[0]; self.rtl(paragraph, "center")
            self.font(paragraph.add_run(header_text), 9, False, MUTED)
            ppr = paragraph._p.get_or_add_pPr(); border = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "C9D4EA")):
                bottom.set(qn(f"w:{k}"), v)
            border.append(bottom); ppr.append(border)
        if footer_numbers:
            paragraph = section.footer.paragraphs[0]; self.rtl(paragraph, "center")
            self.field(paragraph, "PAGE", "۱", 10, MUTED)
        return section

    def blank(self, count=1, size=12):
        for _ in range(count):
            paragraph = self.doc.add_paragraph(); self.spacing(paragraph, 0, 0, 1.0)
            self.font(paragraph.add_run(" "), size)

    def title_page(self, cover):
        self.para("به نام خدا", 11, False, MUTED, "center", 0, 14)
        self.para("دانشگاه آزاد اسلامی", 22 if cover else 18, True, NAVY, "center", 0, 4)
        self.para("رشته‌ی مهندسی کامپیوتر", 15, False, MUTED, "center", 0, 18)
        paragraph = self.doc.add_paragraph(); self.rtl(paragraph, "center"); self.spacing(paragraph, 6, 6, 1.0)
        paragraph.add_run().add_picture(str(FIG_DIR / "logo.png"), width=Cm(3.6 if cover else 2.6))
        self.para("گزارش پروژه", 16, False, "FF6B4A", "center", 14, 6)
        self.para("طراحی و پیاده‌سازی فروشگاه اینترنتی «تک‌شاپ»", 25 if cover else 21, True, NAVY, "center", 0, 8, 1.25)
        self.para("با استفاده از Django، Django REST Framework، React و PostgreSQL", 13, False, MUTED, "center", 0, 40 if cover else 26)
        rows = [("استاد مربوطه", "جناب آقای سجاد پیراهش"), ("تهیه‌کننده", "محمد مهدی پولادرگ"), ("رشته", "مهندسی کامپیوتر"), ("دانشگاه", "دانشگاه آزاد اسلامی"), ("زمان", "مهر ۱۴۰۵")]
        table = self.doc.add_table(rows=len(rows), cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual")); self.borders(table, "C9D4EA", 4)
        for index, (label, value) in enumerate(rows):
            for c, (text, bold, width) in enumerate(((label, True, 4.5), (value, False, 8.5))):
                cell = table.cell(index, c); cell.width = Cm(width); cell.text = ""
                self.shade(cell, NAVY_PALE if c == 0 else "FFFFFF")
                paragraph = cell.paragraphs[0]; self.rtl(paragraph, "center"); self.spacing(paragraph, 3, 3, 1.2)
                self.font(paragraph.add_run(text), 12.5, bold, NAVY if bold else INK)

    def simple_page(self, title, paragraphs, keywords=None, latin=False):
        paragraph = self.doc.add_paragraph(); self.rtl(paragraph, "center" if not latin else "center")
        self.spacing(paragraph, 0, 14, 1.2)
        run = paragraph.add_run(title); self.font(run, self.cfg["h1"], True, NAVY)
        if latin:
            self.latin_paragraphs(paragraphs, keywords)
            return
        for text in paragraphs:
            self.para(text, after=8)
        if keywords:
            self.para(f"کلیدواژه‌ها: {keywords}", self.cfg["body"] - 0.5, True, NAVY, "both", 6, 4)

    def latin_paragraphs(self, paragraphs, keywords):
        for text in paragraphs:
            paragraph = self.doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; self.spacing(paragraph, 0, 8, 1.3)
            run = paragraph.add_run(text); run.font.name = "Times New Roman"; run.font.size = Pt(self.cfg["body"] - (0 if self.flavour == "word" else 0.5))
        paragraph = self.doc.add_paragraph(); self.spacing(paragraph, 6, 4, 1.2)
        run = paragraph.add_run("Keywords: "); run.bold = True; run.font.name = "Times New Roman"; run.font.size = Pt(self.cfg["body"] - 1)
        run = paragraph.add_run(keywords); run.font.name = "Times New Roman"; run.font.size = Pt(self.cfg["body"] - 1)

    def contents_page(self):
        paragraph = self.doc.add_paragraph(); self.rtl(paragraph, "center"); self.spacing(paragraph, 0, 12, 1.2)
        self.font(paragraph.add_run("فهرست مطالب"), self.cfg["h1"], True, NAVY)
        if self.flavour == "word":
            paragraph = self.doc.add_paragraph(); self.rtl(paragraph)
            self.field(paragraph, 'TOC \\o "1-2" \\h \\z \\u', "برای ساخت فهرست مطالب: روی این خط راست‌کلیک کنید و Update Field را بزنید.", self.cfg["body"], MUTED)
            return
        table = self.doc.add_table(rows=len(self.toc_entries), cols=2)
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual")); layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); table._tbl.tblPr.append(layout)
        for index, (level, text) in enumerate(self.toc_entries):
            title_cell, page_cell = table.cell(index, 0), table.cell(index, 1)
            title_cell.width, page_cell.width = Cm(14.0), Cm(2.0)
            for cell, value, align in ((title_cell, text, "both"), (page_cell, fa(self.toc_pages.get(text, "")), "left")):
                cell.text = ""; paragraph = cell.paragraphs[0]; self.rtl(paragraph, align); self.spacing(paragraph, 1, 1, 1.1)
                if align == "left":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.right_indent = Cm(0.0 if level == 1 else 0.8)
                self.font(paragraph.add_run(value), self.cfg["body"] - (0.5 if level == 1 else 1.5), level == 1, NAVY if level == 1 else INK)
            tcpr = title_cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders"); b = OxmlElement("w:bottom")
            for k, v in (("val", "dotted"), ("sz", "4"), ("color", "9DB4DE")):
                b.set(qn(f"w:{k}"), v)
            borders.append(b); tcpr.append(borders)
            tcpr = page_cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders"); b = OxmlElement("w:bottom")
            for k, v in (("val", "dotted"), ("sz", "4"), ("color", "9DB4DE")):
                b.set(qn(f"w:{k}"), v)
            borders.append(b); tcpr.append(borders)

    def lists_page(self, title, entries, kind):
        paragraph = self.doc.add_paragraph(); self.rtl(paragraph, "center"); self.spacing(paragraph, 14, 8, 1.2)
        self.font(paragraph.add_run(title), self.cfg["h2"] + 2, True, NAVY)
        table = self.doc.add_table(rows=max(1, len(entries)), cols=2)
        table._tbl.tblPr.append(OxmlElement("w:bidiVisual")); layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); table._tbl.tblPr.append(layout)
        for index, (number, text) in enumerate(entries):
            title_cell, page_cell = table.cell(index, 0), table.cell(index, 1)
            title_cell.width, page_cell.width = Cm(14.0), Cm(2.0)
            pages = fa(self.list_pages.get((kind, number), "")) if self.flavour == "pdf" else ""
            for cell, value in ((title_cell, f"{kind} {number}: {text}"), (page_cell, pages)):
                cell.text = ""; paragraph = cell.paragraphs[0]; self.rtl(paragraph, "both"); self.spacing(paragraph, 0.5, 0.5, 1.05)
                if cell is page_cell:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self.font(paragraph.add_run(value), self.cfg["body"] - 2)

    # ---------------------------------------------------------------- whole document
    def build(self, path, toc_entries=None, figures=None, tables=None):
        self.toc_entries = toc_entries or [(1, "—")]
        figures_list, tables_list = figures or [("۰", "—")], tables or [("۰", "—")]
        # cover (no header/footer)
        self.title_page(True)
        # title page and front matter
        self.new_section()
        self.title_page(False)
        self.doc.add_page_break()
        self.simple_page("تشکر و قدردانی", THANKS)
        self.doc.add_page_break()
        self.simple_page("چکیده", ABSTRACT_FA, KEYWORDS_FA)
        self.doc.add_page_break()
        self.simple_page("Abstract", ABSTRACT_EN, KEYWORDS_EN, latin=True)
        self.doc.add_page_break()
        self.contents_page()
        self.doc.add_page_break()
        self.lists_page("فهرست شکل‌ها", figures_list, "شکل")
        self.lists_page("فهرست جدول‌ها", tables_list, "جدول")
        # body: numbering restarts at 1
        self.new_section(restart=1, header_text="گزارش پروژه: طراحی و پیاده‌سازی فروشگاه اینترنتی تک‌شاپ", footer_numbers=True)
        self.front_headings = len(self.headings)
        for chapter in (CH1, CH2, CH3, CH4, CH5, CH6, CH7, REFS, APPENDIX):
            self.blocks(chapter)
        if self.flavour == "word":
            settings = self.doc.settings.element
            update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
        self.doc.core_properties.title = "گزارش پروژه: طراحی و پیاده‌سازی فروشگاه اینترنتی تک‌شاپ"
        self.doc.core_properties.author = "محمد مهدی پولادرگ"
        self.doc.save(path)


def build(flavour, path, toc_pages=None, list_pages=None, toc_entries=None, figures=None, tables=None, overrides=None):
    report = Report(flavour, toc_pages, list_pages, overrides)
    report.build(path, toc_entries, figures, tables)
    return report


if __name__ == "__main__":
    flavour, out = sys.argv[1], sys.argv[2]
    build(flavour, out)
    print("saved", out)
